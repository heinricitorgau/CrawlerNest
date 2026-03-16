from __future__ import annotations

import csv
import json
import logging
from abc import ABC, abstractmethod
from pathlib import Path
from typing import List, Optional

from models import University
from utils import truncate_text, format_score

logger = logging.getLogger("UniversityCrawler")


class DataExporter(ABC):
    @abstractmethod
    def export(self, universities: List[University], output_path: Optional[str] = None) -> None:
        pass


class ConsoleExporter(DataExporter):
    def __init__(self, width: int = 150):
        self.width = width

    def export(self, universities: List[University], output_path: Optional[str] = None) -> None:
        if not universities:
            print("No data to display")
            return

        # QS-style short abbreviations for metric column headers
        target_metrics = [
            "Overall Score",
            "Academic Reputation",
            "Employer Reputation",
            "Faculty Student Ratio",
            "Citations per Faculty",
            "International Faculty Ratio",
            "International Student Ratio",
            "International Research Network",
            "Employment Outcomes",
            "Sustainability",
            # Subject specific
            "H-index Citations",
            "Citations per Paper",
            # MBA specific
            "Thought Leadership",
            "Return On Investment",
            "Entrepreneurship & Alumni Outcomes",
            "Employability",
            "Diversity",
            # Business Masters specific
            "Value for Money",
            "Alumni Outcomes",
            # City Rankings specific
            "Student View",
            "Student Mix",
            "Employer Activity",
            "Desirability",
            "Affordability",
            "Rankings",
        ]

        header_aliases = {
            "academic reputation":        "AR",
            "employer reputation":        "ER",
            "faculty student ratio":      "FSR",
            "citations per faculty":      "CPF",
            "international faculty ratio":"IFR",
            "international student ratio":"ISR",
            "international research network": "IRN",
            "employment outcomes":        "EO",
            "sustainability score":       "SUS",
            "sustainability":             "SUS",
            "overall score":              "Score",
            "h-index citations":          "H-index",
            "citations per paper":        "CPP",
            "thought leadership":         "TL",
            "return on investment":       "ROI",
            "entrepreneurship & alumni outcomes": "EAO",
            "employability":              "EMP",
            "diversity":                  "DIV",
            "value for money":            "VFM",
            "alumni outcomes":            "ALU",
            # City Rankings specific
            "student view":               "STV",
            "student mix":                "STM",
            "employer activity":          "EMA",
            "desirability":               "DES",
            "affordability":              "AFF",
            "rankings":                   "RNK",
        }

        def metric_header(name: str) -> str:
            return header_aliases.get(name.strip().lower(), name.strip())

        def trim_score(v: str) -> str:
            """Remove trailing .0 to match QS compact display (e.g. 85.0 -> 85)."""
            try:
                f = float(v)
                if f == int(f):
                    return str(int(f))
                return f"{f:.1f}"
            except (ValueError, TypeError):
                return v

        def metric_value(table_metrics: dict, target_name: str) -> str:
            if not table_metrics:
                return "N/A"
            norm_target = target_name.strip().lower()
            for k, v in table_metrics.items():
                if str(k).strip().lower() == norm_target:
                    return trim_score(str(v))
            for k, v in table_metrics.items():
                kk = str(k).strip().lower()
                if norm_target in kk or kk in norm_target:
                    return trim_score(str(v))
            return "N/A"

        # Detect tied ranks to prefix them with "="
        rank_values = [str(u.rank).lstrip("=").strip() for u in universities]
        rank_counts: dict[str, int] = {}
        for rv in rank_values:
            rank_counts[rv] = rank_counts.get(rv, 0) + 1

        def fmt_rank(raw: str) -> str:
            rv = str(raw).lstrip("=").strip()
            return f"={rv}" if rank_counts.get(rv, 0) > 1 else rv

        rank_w   = max(8,  min(10, max(len(fmt_rank(str(u.rank))) for u in universities) + 2))
        uni_w    = max(30, min(45, max(len(str(u.name))           for u in universities) + 2))
        country_w = max(14, min(22, max(len(str(u.country))       for u in universities) + 2))

        metric_widths: dict[str, int] = {}
        active_metrics: List[str] = []
        for h in target_metrics:
            values = [metric_value(u.table_metrics or {}, h) for u in universities]
            if any(v != "N/A" for v in values):
                active_metrics.append(h)
                short = metric_header(h)
                max_val_len = max(len(v) for v in values)
                metric_widths[h] = max(len(short) + 2, min(10, max_val_len + 2))

        def cell(text: str, width: int) -> str:
            return f"{truncate_text(str(text), width - 1):<{width}}"

        # ── Column widths list (rank, name, country, metrics…) ──
        col_widths = [rank_w, uni_w, country_w] + [metric_widths[h] for h in active_metrics]

        def hline(left: str, mid: str, right: str, fill: str = "─") -> str:
            # Each segment is fill*(width+2) to match " cell " padding
            segs = [fill * (w + 2) for w in col_widths]
            return left + mid.join(segs) + right

        def render_row(parts: List[str]) -> str:
            # Each cell is surrounded by a single space; columns separated by │
            return "│" + "".join(f" {p} │" for p in parts)

        # ── Header row ──
        header_cells = [
            cell("Rank",       rank_w),
            cell("University", uni_w),
            cell("Country",    country_w),
        ] + [cell(metric_header(h), metric_widths[h]) for h in active_metrics]

        # ── Print table ──
        print()
        print(hline("┌", "┬", "┐"))
        print(render_row(header_cells))
        print(hline("├", "┼", "┤"))
        for uni in universities:
            row_parts = [
                cell(fmt_rank(str(uni.rank)), rank_w),
                cell(str(uni.name), uni_w),
                cell(str(uni.country), country_w),
            ]
            for h in active_metrics:
                v = metric_value(uni.table_metrics or {}, h)
                row_parts.append(cell(v, metric_widths[h]))
            print(render_row(row_parts))

        print(hline("└", "┴", "┘"))
        n = len(universities)
        print(f"  Showing 1–{n} of {n} universities\n")

    def print_metric_ranking(
        self,
        universities: List[University],
        metric_choice: str,
        *,
        sort_ascending: bool = False,
    ) -> None:

        if not universities:
            print("No data to display.")
            return
        if all(
            uni.requirements.gmat is None
            and uni.requirements.gre is None
            and uni.requirements.gpa is None
            and uni.requirements.ielts is None
            and uni.requirements.toefl is None
            and uni.requirements.duolingo is None
            and uni.requirements.overall_score is None
            for uni in universities
        ):
            print("This ranking type does not include admission metrics (GMAT/GRE/GPA/IELTS/TOEFL/DET).")
            return

        metric_map = {
            "gmat":    "gmat",
            "gre":     "gre",
            "gpa":     "gpa",
            "ielts":   "ielts",
            "toefl":   "toefl",
            "det":     "duolingo",
            "duolingo":"duolingo",
            "overall": "overall_score",
        }

        if metric_choice not in metric_map:
            print(f"Invalid metric '{metric_choice}'.")
            return

        attr_name    = metric_map[metric_choice]
        metric_label = metric_choice.upper() if metric_choice != "overall" else "Overall Score"

        def key(u: University):
            v = getattr(u.requirements, attr_name)
            if v is None:
                return (1, 0.0)
            try:
                n = float(v)
            except (TypeError, ValueError):
                return (1, 0.0)
            return (0, n if sort_ascending else -n)

        ranked = sorted(universities, key=key)

        dir_text = "lowest → highest" if sort_ascending else "highest → lowest"
        sep = "─" * 110
        print(f"\n{sep}")
        print(f"  {metric_label} Ranking  ({dir_text},  N/A at bottom)")
        print(sep)
        print(f"  {'#':<6}  {metric_label:<14}  {'QS Rank':<10}  {'Country':<22}  University")
        print(sep)

        rank_count = 0
        for uni in ranked:
            val = getattr(uni.requirements, attr_name)
            if val is None:
                rank_display = "–"
                display_val  = "N/A"
            else:
                rank_count  += 1
                rank_display = str(rank_count)
                display_val  = uni.requirements.get_display_value(attr_name)

            print(f"  {rank_display:<6}  {display_val:<14}  {uni.rank:<10}  {uni.country:<22}  {uni.name}")

        print(sep + "\n")


class CSVExporter(DataExporter):
    def __init__(self, **kwargs):
        pass

    def export(self, universities: List[University], output_path: Optional[str] = None) -> None:
        if not output_path:
            output_path = "universities.csv"

        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)

        try:
            with open(path, "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f)
                writer.writerow(
                    [
                        "QS Rank",
                        "University",
                        "Country",
                        "GMAT",
                        "GRE",
                        "GPA",
                        "IELTS",
                        "TOEFL",
                        "Duolingo",
                        "Overall Score",
                        "URL",
                    ]
                )

                for uni in universities:
                    req = uni.requirements
                    writer.writerow(
                        [
                            uni.rank,
                            uni.name,
                            uni.country,
                            format_score(req.gmat),
                            format_score(req.gre),
                            format_score(req.gpa),
                            format_score(req.ielts),
                            format_score(req.toefl),
                            format_score(req.duolingo),
                            format_score(req.overall_score, decimals=1),
                            uni.path,
                        ]
                    )

            logger.info(f"CSV exported to {path}")
            print(f"✓ Data exported to: {path}")

        except Exception as e:
            logger.error(f"Failed to export CSV: {e}")
            print(f"✗ Failed to export CSV: {e}")


class ExcelExporter(DataExporter):
    def __init__(self, **kwargs):
        pass

    def export(self, universities: List[University], output_path: Optional[str] = None) -> None:
        try:
            import openpyxl
            from openpyxl.utils import get_column_letter
            from openpyxl.cell.cell import MergedCell
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            logger.error("openpyxl not installed. Install with: pip install openpyxl")
            print("✗ openpyxl not installed. Install with: pip install openpyxl")
            return

        if not output_path:
            output_path = "universities.xlsx"

        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            assert ws is not None
            ws.title = "University Rankings"

            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")

            headers = [
                "QS Rank",
                "University",
                "Country",
                "GMAT",
                "GRE",
                "GPA",
                "IELTS",
                "TOEFL",
                "Duolingo",
                "Overall Score",
                "URL",
            ]

            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                if isinstance(cell, MergedCell):
                    continue
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")

            for row_idx, uni in enumerate(universities, 2):
                req = uni.requirements
                ws.cell(row=row_idx, column=1, value=uni.rank)
                ws.cell(row=row_idx, column=2, value=uni.name)
                ws.cell(row=row_idx, column=3, value=uni.country)
                ws.cell(row=row_idx, column=4, value=format_score(req.gmat))
                ws.cell(row=row_idx, column=5, value=format_score(req.gre))
                ws.cell(row=row_idx, column=6, value=format_score(req.gpa))
                ws.cell(row=row_idx, column=7, value=format_score(req.ielts))
                ws.cell(row=row_idx, column=8, value=format_score(req.toefl))
                ws.cell(row=row_idx, column=9, value=format_score(req.duolingo))
                ws.cell(row=row_idx, column=10, value=format_score(req.overall_score, decimals=1))
                ws.cell(row=row_idx, column=11, value=uni.path)

            column_widths = [10, 40, 25, 10, 10, 10, 10, 10, 12, 12, 50]
            for col, width in enumerate(column_widths, 1):
                col_letter = get_column_letter(col)
                ws.column_dimensions[col_letter].width = width

            wb.save(path)
            logger.info(f"Excel exported to {path}")
            print(f"✓ Data exported to: {path}")

        except Exception as e:
            logger.error(f"Failed to export Excel: {e}")
            print(f"✗ Failed to export Excel: {e}")


class JSONExporter(DataExporter):
    def __init__(self, **kwargs):
        pass

    def export(self, universities: List[University], output_path: Optional[str] = None) -> None:
        if not output_path:
            output_path = "universities.json"

        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)

        try:
            data = {"total": len(universities), "universities": [u.to_dict() for u in universities]}
            with open(path, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
            logger.info(f"JSON exported to {path}")
            print(f"✓ Data exported to: {path}")
        except Exception as e:
            logger.error(f"Failed to export JSON: {e}")
            print(f"✗ Failed to export JSON: {e}")


def get_exporter(format_type: str, **kwargs) -> DataExporter:
    exporters = {"console": ConsoleExporter, "csv": CSVExporter, "excel": ExcelExporter, "json": JSONExporter}
    if format_type not in exporters:
        logger.warning(f"Unknown format {format_type}, defaulting to console")
        format_type = "console"
    return exporters[format_type](**kwargs)
